"""学习功能的一组纯函数：复盘拼题 / 转写上下文 / 专注文案 / Anki 导出 / 找历史库。

**不带 Qt** —— 全是「给数据、回数据」的东西，测试不用起界面。
拼提示词的口径都在这里，界面层（bar/app）只管把数据搬进搬出。
"""
from __future__ import annotations

import csv
import datetime
import io
import pathlib

from learnsys import knowledge

REVIEW_DISPLAY = "复盘：我刚才学了什么？"      # 上屏显示的那句（真正发给 AI 的题在 study 里拼）

_TRANSCRIPT_CAP = 4000          # 复盘时给 AI 的转写正文上限（字符）
_CONTEXT_CAP = 1200             # 追问时附带的转写上下文上限 —— 太长撑 token、也没那么相关
_CONTEXT_MINUTES = 12           # 追问只带最近这么多分钟的转写（不够就有多少带多少）


def day_dbs(data_root: pathlib.Path) -> list[tuple[datetime.date, pathlib.Path]]:
    """数据根目录下所有「有库的日子」，新→旧。名字是 2026.10.6 这种 ⇒ 按日期解析排序，别按字符串。"""
    out: list[tuple[datetime.date, pathlib.Path]] = []
    if not data_root.is_dir():
        return out
    for db in data_root.glob("*/learnsys.db"):
        parts = db.parent.name.split(".")
        try:
            day = datetime.date(int(parts[0]), int(parts[1]), int(parts[2]))
        except (ValueError, IndexError):
            continue
        out.append((day, db))
    out.sort(key=lambda pair: pair[0], reverse=True)
    return out


def storage_dbs(data_root: pathlib.Path) -> list[tuple[datetime.date, pathlib.Path]]:
    """当前该读哪个库：v2 统一库存在就只读它；否则退回旧按天库。

    迁移后旧按天库仍在（只读备份），再一起扫会把同一条问答数两遍。
    """
    unified = data_root / "learning.db"
    if unified.exists():
        return [(datetime.date.today(), unified)]
    return day_dbs(data_root)


def _transcript_ts_text(row) -> tuple[str, str]:
    """转写行的时间/正文，兼容两种形状：search_transcripts 的 4 列（含 session_id）
    和测试/手工传的 3 列 —— 4 列时正文在 [3]，3 列时在 [2]（真踩过：拿 [2] 当正文，
    摘要全是 ts_end，AI 看到「转写是空的」）。"""
    if len(row) >= 4:
        return str(row[1]), (row[3] or "").strip()
    return str(row[0]), (row[2] or "").strip()


def transcripts_digest(rows) -> str:
    """转写行 → 一段可读正文（带时刻，去空行）。"""
    pieces = []
    for row in rows:
        ts, text = _transcript_ts_text(row)
        if text:
            pieces.append(f"[{ts[11:16] if len(ts) >= 16 else ts}] {text}")
    return "\n".join(pieces)


def build_review_question(summary_lines: list[str], transcript_text: str,
                          frame_text: str = "") -> str:
    """复盘按钮发给 AI 的题：窗口摘要 + 转写正文 + 明确的输出要求。"""
    transcript_text = (transcript_text or "").strip()[-_TRANSCRIPT_CAP:]
    parts = [
        "以下是我今天在电脑上学习的自动记录（看过哪些窗口的摘要 + 老师讲话的声音转写）。",
        "请帮我复盘：①我主要在学什么；②列出最多 5 个值得记住的要点；",
        "③指出哪里看起来没学进去（比如长时间停在无关窗口）；④给我一个 10 分钟以内的复习建议。",
        "用纯文本，简洁，别复述记录本身。",
        "",
        "【窗口摘要】",
    ]
    parts += summary_lines or ["（这段没有窗口记录）"]
    if transcript_text:
        parts += ["", "【声音转写】", transcript_text]
    else:
        parts += ["", "【声音转写】（这段没有转写）"]
    return "\n".join(parts)


def transcript_context(rows, now: datetime.datetime | None = None) -> str:
    """追问要附带的课上上下文：最近 `_CONTEXT_MINUTES` 分钟的转写，超长截尾。

    rows 时间升序，形状兼容 3 列 / 4 列（见 _transcript_ts_text）。没有就回空串。
    """
    if not rows:
        return ""
    now = now or datetime.datetime.now()
    floor = now.timestamp() - _CONTEXT_MINUTES * 60
    pieces = []
    for row in rows:
        ts, text = _transcript_ts_text(row)
        try:
            moment = datetime.datetime.strptime(ts[:19], "%Y-%m-%d %H:%M:%S")
        except ValueError:
            continue
        if moment.timestamp() >= floor and text:
            pieces.append(text)
    if not pieces:
        return ""
    return "（以下是刚录到的课堂转写片段，供参考回答本题）：" + "\n".join(pieces)[-_CONTEXT_CAP:]


def focus_line(stats: dict) -> str:
    """专注统计 → 一句给人看的文案。"""
    minutes = float(stats.get("分钟") or 0)
    line = f"今天记录了 {minutes:.0f} 分钟（{minutes / 60:.1f} 小时），切了 {stats.get('切换', 0)} 次窗口"
    apps = stats.get("程序") or []
    if apps:
        pretty = "、".join(f"{name} {mins:.0f} 分" for name, mins in apps[:3])
        line += f"。最久停在：{pretty}"
    return line


_WRONG_SOURCE_LABEL = {
    "ask": "问答收藏", "quiz": "测验错题", "manual": "手动收录", "cet6": "六级断点",
}


def wrongbook_csv_text(items: list[dict]) -> str:
    """错题条目 → Anki 可导入的 CSV 文本（不含 BOM，调用方决定）。

    正面只放题目；背面把正确答案和解析拼起来，避免自测时被提示。
    """
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(["正面（问题）", "背面（答案）", "我的作答", "来源",
                     "资料", "知识点", "收录时间"])
    for item in items:
        question = str(item.get("question") or "").strip()
        answer = str(item.get("answer") or "").strip()
        note = str(item.get("note") or "").strip()
        if not question:
            continue
        if note and note not in answer:
            answer = (answer + "\n\n" + note).strip()
        writer.writerow([
            question, answer or "（还没有答案）",
            str(item.get("my_answer") or "").strip(),
            _WRONG_SOURCE_LABEL.get(str(item.get("source") or ""), "错题"),
            str(item.get("material_title") or "").strip(),
            str(item.get("tags") or "").strip(),
            str(item.get("ts") or "")[:16],
        ])
    return out.getvalue()


def export_wrongbook_csv(items: list[dict], path: pathlib.Path) -> int:
    """写 utf-8-sig（Excel 双击不乱码，Anki 也认）；返回写入条数。"""
    text = wrongbook_csv_text(items)
    pathlib.Path(path).write_text("﻿" + text, encoding="utf-8")
    # 用 csv 读回再数：答案里的换行在 CSV 里是同一行记录，不能按 splitlines 数。
    rows = list(csv.reader(io.StringIO(text)))
    return max(0, len(rows) - 1)


def export_anki_csv(rows, path: pathlib.Path) -> int:
    """旧调用壳：问答行 [(ts, q, a, kind, image, id), ...] → 统一错题 CSV。"""
    items = [
        {"question": row[1], "answer": row[2], "ts": row[0], "source": "ask"}
        for row in rows
    ]
    return export_wrongbook_csv(items, path)


# ---- 每日报告 / 测验 / 笔记导出 ----

def build_daily_report(day_name: str, summary_lines: list[str], asks: list,
                       transcript_text: str, starred: int, frames: int,
                       ai_summary: str = "") -> str:
    """把一天的数据拼成 Markdown 报告。ai_summary 是 AI 总结那段（可能为空）。"""
    lines = [f"# 学习报告 · {day_name}", ""]
    if ai_summary:
        lines += ["## 要点", "", ai_summary, ""]
    lines += ["## 记录概览", ""]
    lines += (summary_lines or ["（今天没有开过录制）"])
    lines += ["", f"问答 {len(asks)} 条 ｜ 错题本新增 {starred} 条 ｜ 屏幕快照 {frames} 张", ""]
    if asks:
        lines += ["## 今天问过的", ""]
        for row in asks:
            q = (row[1] or "").strip().replace("\n", " ")
            a = (row[2] or "").strip().replace("\n", " ")
            stamp = str(row[0])[11:16]
            lines.append(f"- **{stamp}** {q} —— {a[:120]}{'…' if len(a) > 120 else ''}")
        lines.append("")
    if transcript_text:
        lines += ["## 课堂转写", "", transcript_text[-_TRANSCRIPT_CAP:], ""]
    return "\n".join(lines)


def build_quiz_question(summary_lines: list[str], asks_preview: list[str],
                        transcript_text: str) -> str:
    """测验按钮发给 AI 的题：基于今天的内容出 5 道自测题，逐题批改。"""
    parts = [
        "以下是我今天学习的记录（窗口摘要 + 问过的问题 + 课堂转写）。",
        "请基于这些内容出 **5 道自测题** 考我，要求：",
        "1. 只出题，先不要给答案；**一次只出一题**，每道题以「【第N题】」开头；",
        "2. 题目从记录里的知识点来，别超纲；",
        "3. 有简答也有选择；4. 我答一题你批一题，答错讲清楚错在哪；",
        "5. 只要我答错（含部分作答），先调用 save_wrongbook 工具："
        "question=题目原文（别改写，不要带【第N题】前缀和【知识点】行）、"
        "correct_answer=正确答案或关键解析、my_answer=我的原话、"
        "tags=本题【知识点】里的 1–3 个类目；然后再批改、出下一题。"
        "答对不要调用，同一题只调用一次。",
        "",
        knowledge.prompt_rules(),
        "",
        "【窗口摘要】",
    ]
    parts += summary_lines or ["（今天没有开过录制）"]
    if asks_preview:
        parts += ["", "【今天问过的】"]
        parts += asks_preview
    if transcript_text:
        parts += ["", "【课堂转写】", transcript_text[-_TRANSCRIPT_CAP:]]
    return "\n".join(parts)


def _paged_preview(text: str, cap: int) -> tuple[str, bool]:
    """按页截取正文：PDF 的 `\\f` 换成「--- 第 N 页 ---」，超出 cap 截断。"""
    text = str(text or "")
    pages = text.split("\f") if "\f" in text else [text]
    pieces, used = [], 0
    truncated = False
    for index, page in enumerate(pages, 1):
        chunk = (f"--- 第 {index} 页 ---\n" if len(pages) > 1 else "") + page
        left = cap - used
        if left <= 0:
            truncated = True
            break
        if len(chunk) > left:
            pieces.append(chunk[:left])
            truncated = True
            break
        pieces.append(chunk)
        used += len(chunk)
    return "\n\n".join(pieces).strip(), truncated


def build_material_quiz_question(material: dict, cap: int = 12000) -> str:
    """从一份资料出题的提示词：只基于资料，一次一题；不够就 read_material 继续读。"""
    title = str(material.get("title") or "资料")
    body = str(material.get("body") or "")
    preview, truncated = _paged_preview(body, cap)
    size = f"{material.get('chars') or len(body)} 字"
    if material.get("pages"):
        size += f" / {material.get('pages')} 页"
    parts = [
        f"以下是我资料库里的《{title}》（资料 id={material.get('id')}，{size}）。",
        "请**只基于这份资料**出 **5 道自测题** 考我，要求：",
        "1. 只出题，先不要给答案；**一次只出一题**，每道题以「【第N题】」开头；",
        "2. 覆盖资料里的不同章节，优先定义、计算、易错点，别出资料里没有的内容；",
        "3. 有简答也有选择；我答一题你批一题；",
        "4. 如果给出的正文不够，调用 `read_material(material_id="
        f"{material.get('id')})` 继续读；找关键词也可以 `search_materials`；",
        "5. 只要我答错（含部分作答），先调用 save_wrongbook："
        "question=题目原文（不要带【第N题】前缀和【知识点】行）、"
        "correct_answer=正确答案或关键解析、my_answer=我的原话、"
        "tags=本题【知识点】里的 1–3 个类目；"
        "答对不要调用，同一题只调用一次。",
        "",
        knowledge.prompt_rules(),
        "",
        "【资料正文】",
        preview or "（这份资料还没有正文）",
    ]
    if truncated:
        parts += ["", f"（正文只给了前 {cap} 字，其余请用 read_material 继续读。）"]
    return "\n".join(parts)


def export_day_notes_md(asks: list, path: pathlib.Path, title: str) -> int:
    """问答行 → Markdown 笔记（一问一段）。返回写入条数。"""
    count = 0
    out = [f"# {title}", ""]
    for row in asks:
        q = (row[1] or "").strip()
        a = (row[2] or "").strip()
        if not q:
            continue
        stamp = str(row[0])[:16]
        kind = row[3] if len(row) > 3 else "text"
        tag = "【图问】" if kind == "image" else "【文问】"
        out += [f"## {stamp} {tag}{q}", "", a or "（无答案）", ""]
        count += 1
    pathlib.Path(path).write_text("\n".join(out), encoding="utf-8")
    return count


# ---- 文档问答 / 术语表 ----

def extract_document(path: str) -> dict:
    """课件/文档 → 正文 + 页数 + 类型。

    PDF 用换页符 `\\f` 分页、PPTX 用「--- 第 N 页 ---」分页，资料检索靠它算页码；
    DOCX/TXT/MD 没有可靠页边界，页数记 0。
    """
    p = pathlib.Path(path)
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        import pymupdf
        pages = []
        with pymupdf.open(str(p)) as doc:
            for page in doc:
                pages.append(page.get_text())
        return {"kind": "pdf", "text": "\f".join(pages), "pages": len(pages)}
    if suffix == ".docx":
        import docx
        doc = docx.Document(str(p))
        text = "\n".join(par.text for par in doc.paragraphs if par.text.strip())
        return {"kind": "docx", "text": text, "pages": 0}
    if suffix == ".pptx":
        from pptx import Presentation
        prs = Presentation(str(p))
        parts = []
        for index, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    t = shape.text_frame.text.strip()
                    if t:
                        texts.append(t)
            if texts:
                parts.append(f"--- 第 {index} 页 ---\n" + "\n".join(texts))
        return {"kind": "pptx", "text": "\n\n".join(parts), "pages": len(prs.slides)}
    if suffix in (".txt", ".md"):
        return {"kind": suffix[1:],
                "text": p.read_text(encoding="utf-8", errors="replace"), "pages": 0}
    raise ValueError(f"还不支持 {suffix} 文档（先支持 PDF/DOCX/PPTX/TXT/MD）")


def extract_doc_text(path: str) -> str:
    """旧接口：只取正文（一次性附件用）。新代码优先用 `extract_document`。

    一次性附件不需要页码，把换页符换成换行，别把控制字符塞进提示词。
    """
    return extract_document(path)["text"].replace("\f", "\n")

def build_glossary_question(transcript_text: str, asks_preview: list[str]) -> str:
    """术语表：AI 从今天的转写+问答里提取学科术语。"""
    parts = [
        "以下是我今天学习的记录。请从中提取 **学科术语/关键词**，整理成一张术语表：",
        "每行一条，格式：术语 —— 一句话解释（解释要通俗，20 字以内）；",
        "只收学科相关的词，别收日常用语；按出现顺序排。",
        "",
    ]
    if asks_preview:
        parts += ["【问过的问题】"] + asks_preview + [""]
    if transcript_text:
        parts += ["【课堂转写】", transcript_text[-_TRANSCRIPT_CAP:]]
    return chr(10).join(parts)
